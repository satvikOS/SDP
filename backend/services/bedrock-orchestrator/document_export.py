"""Document export handlers for PDF, PPTX, EPUB, and WORD formats."""

import json
import logging
import os
import io
import base64
from datetime import datetime
from typing import Dict, Any, List

logger = logging.getLogger()
logger.setLevel(os.getenv('LOG_LEVEL', 'INFO'))


def export_to_pdf(event, context):
    """
    Export scenario document to PDF format.

    Request body:
    {
        "scenario": {
            "company_name": str,
            "industry": str,
            "region": str,
            "horizon_years": int,
            "created_at": str,
            "matrix_framework": {...},
            "scenarios": [...]
        }
    }

    Returns:
    {
        "statusCode": 200,
        "headers": {
            "Content-Type": "application/pdf",
            "Content-Disposition": "attachment; filename=..."
        },
        "body": base64_encoded_pdf,
        "isBase64Encoded": true
    }
    """
    try:
        body = json.loads(event['body']) if isinstance(event.get('body'), str) else event.get('body', {})
        scenario_data = body.get('scenario', {})

        if not scenario_data:
            return {
                'statusCode': 400,
                'body': json.dumps({'error': 'Missing scenario data'})
            }

        # Generate PDF using reportlab
        pdf_bytes = generate_pdf(scenario_data)

        # Encode to base64 for API Gateway
        pdf_base64 = base64.b64encode(pdf_bytes).decode('utf-8')

        # Generate filename
        company_name = scenario_data.get('company_name', 'Company').replace(' ', '_')
        filename = f"Strategic_Foresight_{company_name}_{datetime.now().strftime('%Y%m%d')}.pdf"

        return {
            'statusCode': 200,
            'headers': {
                'Content-Type': 'application/pdf',
                'Content-Disposition': f'attachment; filename="{filename}"',
                'Access-Control-Allow-Origin': '*',
                'Access-Control-Allow-Headers': 'Content-Type',
            },
            'body': pdf_base64,
            'isBase64Encoded': True
        }

    except Exception as e:
        logger.error(f"PDF export failed: {str(e)}", exc_info=True)
        return {
            'statusCode': 500,
            'body': json.dumps({'error': f'PDF export failed: {str(e)}'})
        }


def generate_pdf(scenario_data: Dict[str, Any]) -> bytes:
    """Generate PDF document from scenario data using reportlab."""
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.75*inch, bottomMargin=0.75*inch)

    # Container for the 'Flowable' objects
    elements = []

    # Get styles
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=colors.HexColor('#1a202c'),
        spaceAfter=30,
        alignment=TA_CENTER,
        fontName='Helvetica-Bold'
    )

    heading_style = ParagraphStyle(
        'CustomHeading',
        parent=styles['Heading2'],
        fontSize=16,
        textColor=colors.HexColor('#2563eb'),
        spaceBefore=20,
        spaceAfter=12,
        fontName='Helvetica-Bold'
    )

    body_style = ParagraphStyle(
        'CustomBody',
        parent=styles['BodyText'],
        fontSize=11,
        textColor=colors.HexColor('#374151'),
        alignment=TA_JUSTIFY,
        spaceAfter=12,
        leading=14
    )

    # Cover Page
    elements.append(Spacer(1, 2*inch))
    elements.append(Paragraph("STRATEGIC FORESIGHT ANALYSIS", title_style))
    elements.append(Spacer(1, 0.5*inch))

    company_line = f"{scenario_data.get('company_name', 'Company')} | {scenario_data.get('industry', 'Industry')} Sector"
    elements.append(Paragraph(company_line, styles['Heading2']))
    elements.append(Spacer(1, 0.2*inch))

    scope_line = f"{scenario_data.get('region', 'Global')} | {scenario_data.get('horizon_years', 10)}-Year Horizon"
    elements.append(Paragraph(scope_line, styles['Normal']))
    elements.append(Spacer(1, 1*inch))

    elements.append(Paragraph("Prepared for:", styles['Normal']))
    elements.append(Paragraph(f"{scenario_data.get('company_name', 'Company')} Board of Directors", styles['Heading3']))
    elements.append(Spacer(1, 0.5*inch))

    elements.append(Paragraph("Prepared by:", styles['Normal']))
    elements.append(Paragraph("Strategic Foresight Partners", styles['Heading3']))
    elements.append(Spacer(1, 0.5*inch))

    created_date = datetime.fromisoformat(scenario_data.get('created_at', datetime.now().isoformat()).replace('Z', '+00:00'))
    elements.append(Paragraph(created_date.strftime('%B %d, %Y'), styles['Normal']))
    elements.append(Spacer(1, 0.5*inch))

    # Classification marking
    classification = Paragraph(
        '<para align="center" backColor="#fee2e2" borderColor="#dc2626" borderWidth="1" borderPadding="10">'
        '<b>CONFIDENTIAL - BOARD LEVEL ONLY</b></para>',
        styles['Normal']
    )
    elements.append(classification)
    elements.append(PageBreak())

    # Scenario Matrix Framework
    if scenario_data.get('matrix_framework'):
        matrix = scenario_data['matrix_framework']
        elements.append(Paragraph("Scenario Planning Framework", heading_style))

        if matrix.get('axis_x'):
            axis_x = matrix['axis_x']
            elements.append(Paragraph(f"<b>Axis X:</b> {axis_x.get('name', '')}", body_style))
            elements.append(Paragraph(f"{axis_x.get('left_pole', '')} ← → {axis_x.get('right_pole', '')}", body_style))
            elements.append(Paragraph(axis_x.get('description', ''), body_style))
            elements.append(Spacer(1, 0.2*inch))

        if matrix.get('axis_y'):
            axis_y = matrix['axis_y']
            elements.append(Paragraph(f"<b>Axis Y:</b> {axis_y.get('name', '')}", body_style))
            elements.append(Paragraph(f"{axis_y.get('bottom_pole', '')} ← → {axis_y.get('top_pole', '')}", body_style))
            elements.append(Paragraph(axis_y.get('description', ''), body_style))

        elements.append(PageBreak())

    # Scenarios
    scenarios = scenario_data.get('scenarios', [])
    for idx, scenario in enumerate(scenarios, 1):
        # Scenario header
        scenario_title = f"SCENARIO {idx}: {scenario.get('title', 'Untitled')}"
        elements.append(Paragraph(scenario_title, title_style))

        tagline = scenario.get('tagline', '')
        if tagline:
            elements.append(Paragraph(f'<i>{tagline}</i>', styles['Heading3']))
            elements.append(Spacer(1, 0.2*inch))

        # Metadata
        meta_data = [
            ['Quadrant:', scenario.get('quadrant', 'N/A')],
            ['Probability:', f"{scenario.get('probability', 0) * 100:.0f}%"]
        ]
        meta_table = Table(meta_data, colWidths=[1.5*inch, 4*inch])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#e5e7eb')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#374151')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ]))
        elements.append(meta_table)
        elements.append(Spacer(1, 0.3*inch))

        # Core logic
        if scenario.get('core_logic'):
            elements.append(Paragraph("<b>Core Logic</b>", heading_style))
            elements.append(Paragraph(scenario['core_logic'], body_style))
            elements.append(Spacer(1, 0.2*inch))

        # Narrative
        if scenario.get('narrative'):
            elements.append(Paragraph("<b>Strategic Analysis</b>", heading_style))
            narrative_paragraphs = scenario['narrative'].split('\n\n')
            for para in narrative_paragraphs:
                if para.strip():
                    elements.append(Paragraph(para.strip(), body_style))
            elements.append(Spacer(1, 0.3*inch))

        # Signposts
        if scenario.get('signposts') and len(scenario['signposts']) > 0:
            elements.append(Paragraph("<b>Strategic Signposts</b>", heading_style))

            signpost_data = [['Indicator', 'Timeframe', 'Significance']]
            for signpost in scenario['signposts']:
                signpost_data.append([
                    signpost.get('indicator', ''),
                    signpost.get('timeframe', ''),
                    signpost.get('significance', '')
                ])

            signpost_table = Table(signpost_data, colWidths=[2*inch, 1.5*inch, 3*inch])
            signpost_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#3b82f6')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
                ('TOPPADDING', (0, 0), (-1, -1), 8),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ]))
            elements.append(signpost_table)

        # Add page break after each scenario except the last
        if idx < len(scenarios):
            elements.append(PageBreak())

    # Footer with branding
    def add_footer(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 9)
        canvas.setFillColor(colors.grey)
        footer_text = f"{scenario_data.get('company_name', 'Company')} | Strategic Foresight Partners"
        canvas.drawString(inch, 0.5*inch, footer_text)
        canvas.drawRightString(letter[0] - inch, 0.5*inch, f"Page {doc.page}")
        canvas.restoreState()

    # Build PDF
    doc.build(elements, onFirstPage=add_footer, onLaterPages=add_footer)

    pdf_bytes = buffer.getvalue()
    buffer.close()

    return pdf_bytes


def export_to_pptx(event, context):
    """Export scenario document to PowerPoint format."""
    try:
        body = json.loads(event['body']) if isinstance(event.get('body'), str) else event.get('body', {})
        scenario_data = body.get('scenario', {})

        if not scenario_data:
            return {
                'statusCode': 400,
                'body': json.dumps({'error': 'Missing scenario data'})
            }

        # Generate PPTX using python-pptx
        pptx_bytes = generate_pptx(scenario_data)

        # Encode to base64
        pptx_base64 = base64.b64encode(pptx_bytes).decode('utf-8')

        # Generate filename
        company_name = scenario_data.get('company_name', 'Company').replace(' ', '_')
        filename = f"Strategic_Foresight_{company_name}_{datetime.now().strftime('%Y%m%d')}.pptx"

        return {
            'statusCode': 200,
            'headers': {
                'Content-Type': 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
                'Content-Disposition': f'attachment; filename="{filename}"',
                'Access-Control-Allow-Origin': '*',
                'Access-Control-Allow-Headers': 'Content-Type',
            },
            'body': pptx_base64,
            'isBase64Encoded': True
        }

    except Exception as e:
        logger.error(f"PPTX export failed: {str(e)}", exc_info=True)
        return {
            'statusCode': 500,
            'body': json.dumps({'error': f'PPTX export failed: {str(e)}'})
        }


def generate_pptx(scenario_data: Dict[str, Any]) -> bytes:
    """Generate PowerPoint presentation from scenario data."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Title slide
    title_slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(title_slide_layout)

    # Add title
    left = Inches(1)
    top = Inches(2)
    width = Inches(8)
    height = Inches(1)

    title_box = slide.shapes.add_textbox(left, top, width, height)
    title_frame = title_box.text_frame
    title_frame.text = "STRATEGIC FORESIGHT ANALYSIS"
    title_para = title_frame.paragraphs[0]
    title_para.alignment = PP_ALIGN.CENTER
    title_para.font.size = Pt(44)
    title_para.font.bold = True
    title_para.font.color.rgb = RGBColor(26, 32, 44)

    # Company info
    info_box = slide.shapes.add_textbox(Inches(1), Inches(3.5), Inches(8), Inches(0.8))
    info_frame = info_box.text_frame
    info_frame.text = f"{scenario_data.get('company_name', 'Company')} | {scenario_data.get('industry', 'Industry')}"
    info_para = info_frame.paragraphs[0]
    info_para.alignment = PP_ALIGN.CENTER
    info_para.font.size = Pt(24)

    # Classification
    class_box = slide.shapes.add_textbox(Inches(2), Inches(6.5), Inches(6), Inches(0.5))
    class_frame = class_box.text_frame
    class_frame.text = "CONFIDENTIAL - BOARD LEVEL ONLY"
    class_para = class_frame.paragraphs[0]
    class_para.alignment = PP_ALIGN.CENTER
    class_para.font.size = Pt(14)
    class_para.font.bold = True
    class_para.font.color.rgb = RGBColor(220, 38, 38)

    # Add scenario slides
    scenarios = scenario_data.get('scenarios', [])
    for idx, scenario in enumerate(scenarios, 1):
        # Create slide for each scenario
        content_layout = prs.slide_layouts[5]  # Title only layout
        slide = prs.slides.add_slide(content_layout)

        # Title
        title = slide.shapes.title
        title.text = f"Scenario {idx}: {scenario.get('title', 'Untitled')}"
        title.text_frame.paragraphs[0].font.size = Pt(32)
        title.text_frame.paragraphs[0].font.color.rgb = RGBColor(37, 99, 235)

        # Content
        left = Inches(0.5)
        top = Inches(1.5)
        width = Inches(9)
        height = Inches(5)

        content_box = slide.shapes.add_textbox(left, top, width, height)
        text_frame = content_box.text_frame
        text_frame.word_wrap = True

        # Tagline
        if scenario.get('tagline'):
            p = text_frame.paragraphs[0]
            p.text = scenario['tagline']
            p.font.size = Pt(18)
            p.font.italic = True
            p.space_after = Pt(12)

        # Core logic
        if scenario.get('core_logic'):
            p = text_frame.add_paragraph()
            p.text = f"Core Logic: {scenario['core_logic'][:300]}..."  # Truncate for slide
            p.font.size = Pt(14)
            p.space_after = Pt(12)

        # Probability
        p = text_frame.add_paragraph()
        p.text = f"Probability: {scenario.get('probability', 0) * 100:.0f}%"
        p.font.size = Pt(16)
        p.font.bold = True

    # Save to bytes
    buffer = io.BytesIO()
    prs.save(buffer)
    pptx_bytes = buffer.getvalue()
    buffer.close()

    return pptx_bytes


def export_to_word(event, context):
    """Export scenario document to Word format."""
    try:
        body = json.loads(event['body']) if isinstance(event.get('body'), str) else event.get('body', {})
        scenario_data = body.get('scenario', {})

        if not scenario_data:
            return {
                'statusCode': 400,
                'body': json.dumps({'error': 'Missing scenario data'})
            }

        # Generate DOCX using python-docx
        docx_bytes = generate_word(scenario_data)

        # Encode to base64
        docx_base64 = base64.b64encode(docx_bytes).decode('utf-8')

        # Generate filename
        company_name = scenario_data.get('company_name', 'Company').replace(' ', '_')
        filename = f"Strategic_Foresight_{company_name}_{datetime.now().strftime('%Y%m%d')}.docx"

        return {
            'statusCode': 200,
            'headers': {
                'Content-Type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                'Content-Disposition': f'attachment; filename="{filename}"',
                'Access-Control-Allow-Origin': '*',
                'Access-Control-Allow-Headers': 'Content-Type',
            },
            'body': docx_base64,
            'isBase64Encoded': True
        }

    except Exception as e:
        logger.error(f"Word export failed: {str(e)}", exc_info=True)
        return {
            'statusCode': 500,
            'body': json.dumps({'error': f'Word export failed: {str(e)}'})
        }


def generate_word(scenario_data: Dict[str, Any]) -> bytes:
    """Generate Word document from scenario data."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Cover page
    title = doc.add_heading('STRATEGIC FORESIGHT ANALYSIS', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    company_para = doc.add_paragraph()
    company_para.add_run(f"{scenario_data.get('company_name', 'Company')} | {scenario_data.get('industry', 'Industry')}").bold = True
    company_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    prepared_for = doc.add_paragraph()
    prepared_for.add_run("Prepared for: ").bold = True
    prepared_for.add_run(f"{scenario_data.get('company_name', 'Company')} Board of Directors")
    prepared_for.alignment = WD_ALIGN_PARAGRAPH.CENTER

    prepared_by = doc.add_paragraph()
    prepared_by.add_run("Prepared by: ").bold = True
    prepared_by.add_run("Strategic Foresight Partners")
    prepared_by.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    classification = doc.add_paragraph("CONFIDENTIAL - BOARD LEVEL ONLY")
    classification.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = classification.runs[0]
    run.bold = True
    run.font.color.rgb = RGBColor(220, 38, 38)

    doc.add_page_break()

    # Scenario Matrix
    if scenario_data.get('matrix_framework'):
        doc.add_heading('Scenario Planning Framework', 1)
        matrix = scenario_data['matrix_framework']

        if matrix.get('axis_x'):
            axis_x = matrix['axis_x']
            doc.add_heading(f"Axis X: {axis_x.get('name', '')}", 2)
            doc.add_paragraph(f"{axis_x.get('left_pole', '')} ← → {axis_x.get('right_pole', '')}")
            doc.add_paragraph(axis_x.get('description', ''))

        if matrix.get('axis_y'):
            axis_y = matrix['axis_y']
            doc.add_heading(f"Axis Y: {axis_y.get('name', '')}", 2)
            doc.add_paragraph(f"{axis_y.get('bottom_pole', '')} ← → {axis_y.get('top_pole', '')}")
            doc.add_paragraph(axis_y.get('description', ''))

        doc.add_page_break()

    # Scenarios
    scenarios = scenario_data.get('scenarios', [])
    for idx, scenario in enumerate(scenarios, 1):
        doc.add_heading(f"Scenario {idx}: {scenario.get('title', 'Untitled')}", 1)

        if scenario.get('tagline'):
            tagline_para = doc.add_paragraph(scenario['tagline'])
            tagline_para.runs[0].italic = True

        # Metadata
        meta_para = doc.add_paragraph()
        meta_para.add_run(f"Quadrant: ").bold = True
        meta_para.add_run(f"{scenario.get('quadrant', 'N/A')}\n")
        meta_para.add_run(f"Probability: ").bold = True
        meta_para.add_run(f"{scenario.get('probability', 0) * 100:.0f}%")

        # Core logic
        if scenario.get('core_logic'):
            doc.add_heading('Core Logic', 2)
            doc.add_paragraph(scenario['core_logic'])

        # Narrative
        if scenario.get('narrative'):
            doc.add_heading('Strategic Analysis', 2)
            narrative_paragraphs = scenario['narrative'].split('\n\n')
            for para in narrative_paragraphs:
                if para.strip():
                    doc.add_paragraph(para.strip())

        # Signposts
        if scenario.get('signposts') and len(scenario['signposts']) > 0:
            doc.add_heading('Strategic Signposts', 2)
            table = doc.add_table(rows=1 + len(scenario['signposts']), cols=3)
            table.style = 'Light Grid Accent 1'

            # Header
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Indicator'
            hdr_cells[1].text = 'Timeframe'
            hdr_cells[2].text = 'Significance'

            # Data
            for i, signpost in enumerate(scenario['signposts'], 1):
                row_cells = table.rows[i].cells
                row_cells[0].text = signpost.get('indicator', '')
                row_cells[1].text = signpost.get('timeframe', '')
                row_cells[2].text = signpost.get('significance', '')

        if idx < len(scenarios):
            doc.add_page_break()

    # Save to bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()
    buffer.close()

    return docx_bytes


def export_to_epub(event, context):
    """Export scenario document to EPUB format."""
    return {
        'statusCode': 501,
        'body': json.dumps({
            'error': 'EPUB export not yet implemented',
            'message': 'EPUB export will be available in a future update'
        })
    }
